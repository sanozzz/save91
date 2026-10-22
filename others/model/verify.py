"""Recalculate savings_first_model.xlsx for bear/base/bull and assert it
matches engine.py. Exits non-zero on any mismatch or formula error.

Two independent implementations agreeing is the check that the workbook's
formulas compute what the model says they compute.
"""
from __future__ import annotations

import math
import sys
import warnings

import formulas
import openpyxl

import engine
from build_xlsx import OUT

warnings.filterwarnings("ignore")
BOOK = OUT.name
M = engine.MONTHS


def key(sheet, cell):
    return f"'[{BOOK}]{sheet.upper()}'!{cell}"


def main():
    wb = openpyxl.load_workbook(OUT)
    names = {n: wb.defined_names[n].attr_text for n in wb.defined_names}

    def addr(name):
        sheet, cell = names[name].rsplit("!", 1)
        return sheet.strip("'"), cell.replace("$", "")

    coh = wb["Cohort total"]
    label_row = {coh.cell(r, 1).value: r for r in range(1, coh.max_row + 1)}
    cm_row = label_row["Contribution margin"]
    rev_row = label_row["Total revenue"]
    inp = wb["Inputs"]
    input_row = {inp.cell(r, 1).value: r for r in range(1, inp.max_row + 1)}

    xl = formulas.ExcelModel().loads(str(OUT)).finish()
    failures, checked = [], 0
    col_of = {"bear": "C", "base": "D", "bull": "E"}
    # defaults = the founder's plan (dlg_on=1, gold_on=1); also check the idea as briefed and each switch
    cases = [(s, {}) for s in engine.SCENARIOS] + [
        ("base", {"gold_on": 0, "dlg_on": 0}), ("bear", {"gold_on": 0, "dlg_on": 0}), ("base", {"dlg_on": 0}),
        ("base", {"nbfc_on": 1}), ("bull", {"nbfc_on": 1, "dlg_on": 0}),
        # exercise the non-gig input path: different non-gig values in the base column
        ("base", {"ret_m12@ng": 0.31, "daily_amount@ng": 61.0, "ladder_dpd90@ng": 0.031, "cac_mult@ng": 0.9})]
    for s, levers in cases:
        i = engine.SCENARIOS.index(s) + 1
        ins = {key("Inputs", "C2"): i}
        for nm, val in levers.items():
            if nm.endswith("@ng"):   # the non-gig column for this scenario: G / H / I
                ng_col = {"bear": "G", "base": "H", "bull": "I"}[s]
                ins[key("Inputs", f"{ng_col}{input_row[nm[:-3]]}")] = val
                continue
            _, cell = addr(nm)
            ins[key("Inputs", col_of[s] + cell[1:])] = val
        sol = xl.calculate(inputs=ins)
        tag = s + ("[" + ",".join(f"{k}={v}" for k, v in levers.items()) + "]" if levers else "[plan]")

        def val(sheet, cell):
            v = sol[key(sheet, cell)].value
            v = v[0][0] if hasattr(v, "__getitem__") and not isinstance(v, str) else v
            return v

        def by_name(n):
            return val(*addr(n))

        p, rows, out, _ = engine.run(s, dict(levers))
        _, lout = engine.lending_first(engine.load_lf(s), p["disc_rate"])
        prows = engine.path(p, out)

        def check(label, got, want, tol=1e-6):
            nonlocal checked
            checked += 1
            if isinstance(want, str) or want is None:
                ok = (want is None and isinstance(got, str)) or got == want
            else:
                try:
                    ok = math.isclose(float(got), want, rel_tol=tol, abs_tol=1e-6)
                except (TypeError, ValueError):
                    ok = False
            if not ok:
                failures.append(f"[{tag}] {label}: workbook={got!r} engine={want!r}")

        check("CAC", by_name("out_cac"), out["cac"])
        check("payback", by_name("out_payback"), out["payback_months"])
        check("capital per revenue", by_name("out_cap_per_rev"), out["capital_per_rev"])
        for sg in out["seg"]:
            so = out["seg"][sg]
            check(f"{sg} CAC", by_name(f"seg_{sg}_cac"), out["seg_cac"][sg])
            check(f"{sg} LTV60", by_name(f"seg_{sg}_ltv60"), so["ltv60"])
            check(f"{sg} LTV60/CAC", by_name(f"seg_{sg}_ltv60_cac"), so["ltv60_cac"])
            check(f"{sg} payback", by_name(f"seg_{sg}_payback"), so["payback_months"])
            check(f"{sg} rev/active y2", by_name(f"seg_{sg}_rev_pa_y2"), so["y2_rev_per_active"])
            check(f"{sg} still m12", by_name(f"seg_{sg}_still_m12"), so["m12_still"])
        for k in ("ltv36", "ltv60", "ltv36_cac", "ltv60_cac"):
            check(k, by_name(f"out_{k}"), out[k])
        for y in range(1, 6):
            check(f"y{y} rev/active", by_name(f"y{y}_rev_pa"), out[f"y{y}_rev_per_active"])
            check(f"y{y} cm/active", by_name(f"y{y}_cm_pa"), out[f"y{y}_cm_per_active"])
            check(f"y{y} rev/activated", by_name(f"y{y}_rev_act"), out[f"y{y}_rev_per_activated"])
        st = engine.story(p, rows, out)
        for k, v in st.items():
            check(f"story {k}", by_name(k if k.startswith("y") else f"story_{k}"), v)
        for m in range(1, M + 1):
            col = openpyxl.utils.get_column_letter(3 + m)
            check(f"cohort revenue m{m}", val("Cohort total", f"{col}{rev_row}"), rows[m - 1]["revenue"])
            check(f"cohort contribution m{m}", val("Cohort total", f"{col}{cm_row}"), rows[m - 1]["contribution"])
        for k, n in (("cac", "lf_cac"), ("ltv36", "lf_ltv36"), ("ltv60", "lf_ltv60"), ("ltv60_cac", "lf_ltv60_cac"),
                     ("first_loan_contribution", "lf_first_cm"), ("cac_per_good_borrower", "lf_cac_good"),
                     ("capital_per_rev", "lf_cap_per_rev"), ("y1_cm_per_borrower", "lf_y1_cm")):
            check(f"lending-first {k}", by_name(n), lout[k])
        check("lending-first payback", by_name("lf_payback"), lout["payback_months"])
        for r in prows:
            check(f"path rev $M y{r['year']}", by_name(f"path_rev_usd_y{r['year']}"), r["rev_usd_m"])
            check(f"path active y{r['year']}", by_name(f"path_active_y{r['year']}"), r["active"] / 1e6)
        peak = max(0.0, -min(r["cum_cash"] for r in prows)) / 1e7
        check("path peak capital", by_name("path_peak_capital"), peak)

        # no error values anywhere
        errs = [k for k, v in sol.items() if isinstance(getattr(v, "value", None), object)
                and str(getattr(v, "value", "")).startswith(("[[#", "#"))]
        errs = [k for k in errs if any(tok in str(sol[k].value) for tok in ("#VALUE", "#REF", "#NAME", "#DIV", "#N/A", "#NUM"))]
        if errs:
            failures.append(f"[{tag}] {len(errs)} formula errors, e.g. {errs[:5]}")
        print(f"{tag}: CAC {out['cac']:.0f}  LTV60/CAC {out['ltv60_cac']:.2f}  path Y10 ${prows[-1]['rev_usd_m']:.0f}M  peak capital Rs {max(0.0, -min(r['cum_cash'] for r in prows))/1e7:.0f} cr")

    print(f"checked {checked} values across {len(cases)} cases")
    import json
    (engine.TABLES / "verify_result.json").write_text(json.dumps(
        {"checked": checked, "cases": len(cases), "mismatches": len(failures)}, indent=1))
    if failures:
        print(f"{len(failures)} MISMATCHES:")
        for f in failures[:40]:
            print("  " + f)
        sys.exit(1)
    print("workbook == engine for every checked value")


if __name__ == "__main__":
    main()
