"""Workbook tabs for the uncertainty work: the assumption register, the Monte Carlo
results and the one-at-a-time sensitivity on net value. Values are written by mc.py
(the reference engine) from frozen tables; these tabs show them, they don't recompute.
"""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill

T = Path(__file__).resolve().parent.parent / "tables"
F = "Arial"
TITLE = Font(name=F, size=14, bold=True)
BOLD = Font(name=F, size=10, bold=True)
NOTE = Font(name=F, size=9, italic=True, color="555555")
BODY = Font(name=F, size=10)
HEAD_FILL = PatternFill("solid", fgColor="14213D")
HEAD = Font(name=F, size=10, bold=True, color="FFFFFF")
YELLOW = PatternFill("solid", fgColor="FFF2CC")


def read(name):
    p = T / name
    return list(csv.DictReader(open(p))) if p.exists() else []


def num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return v
    return f


def table(ws, r, title, rows, cols, fmts=None, widths=None, highlight=None):
    ws.cell(r, 1, title).font = BOLD
    r += 1
    for j, c in enumerate(cols, 1):
        cell = ws.cell(r, j, c)
        cell.font, cell.fill = HEAD, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        r += 1
        for j, c in enumerate(cols, 1):
            cell = ws.cell(r, j, num(row.get(c, "")))
            cell.font = BODY
            if fmts and c in fmts:
                cell.number_format = fmts[c]
            if highlight and highlight(row):
                cell.fill = YELLOW
    if widths:
        for j, w in enumerate(widths, 1):
            col = ws.cell(1, j).column_letter
            ws.column_dimensions[col].width = max(ws.column_dimensions[col].width or 0, w)
    return r + 2


def build(wb):
    names = {r["id"]: r["name"] for r in read("inputs.csv")}

    # ------------------------------------------------ assumption register
    ws = wb.create_sheet("Assumption register")
    ws["A1"] = "Assumption register: every uncertain input, its distribution, the evidence for it and how much it matters"
    ws["A1"].font = TITLE
    ws["A2"] = ("One row per input per segment (gig / ng = self-employed / all). Distribution = what the Monte Carlo samples. "
                "Status: triangulated = checked against independent sources (see Evidence); provisional = bear/base/bull widened, "
                "no independent check yet (yellow); fixed = founder decision or product design. Variance share = share of the "
                "spread in net value per saver this input explains (Monte Carlo, independent variant).")
    ws["A2"].font = NOTE
    reg = read("assumption_register.csv")
    if not reg:
        drv = {r["input"]: r for r in read("mc_drivers.csv") if r["variant"] == "independent"}
        for d in read("distributions.csv"):
            k = d["input"] if d["segment"] == "all" else f"{d['input']}@{d['segment']}"
            reg.append({"input": d["input"], "segment": d["segment"], "assumption": names.get(d["input"], ""), "dist": d["dist"],
                        "low": d["low"], "mode": d["mode"], "high": d["high"], "status": d["status"], "basis": d["basis"],
                        "evidence": d["evidence"], "variance_share": drv.get(k, {}).get("variance_share", ""),
                        "spearman": drv.get(k, {}).get("spearman", "")})
    cols = [c for c in ("input", "segment", "assumption", "dist", "low", "mode", "high", "status", "verdict", "basis", "evidence",
                        "variance_share", "spearman") if any(c in r for r in reg)]
    table(ws, 4, "Register", reg, cols, fmts={"variance_share": "0.0%", "spearman": "0.00"},
          widths=[18, 8, 44, 11, 10, 10, 10, 13, 14, 50, 30, 10, 9], highlight=lambda r: r.get("status") == "provisional")
    ws.freeze_panes = "D6"

    # ------------------------------------------------ Monte Carlo
    ws = wb.create_sheet("Monte Carlo")
    ws["A1"] = "Monte Carlo: how likely is a profitable saver, given what we know about each input"
    ws["A1"].font = TITLE
    ws["A2"] = ("Written by model/mc.py from tables/distributions.csv: 10,000 draws, fixed seed 20261002, each run through the "
                "reference engine (which verify.py checks against this workbook's formulas). 'independent' = inputs drawn "
                "independently except the retention curve, which moves as one; 'correlated' = an engagement factor links "
                "retention, amounts and take-up (rho 0.5). Net = LTV (60 months, discounted) minus CAC, per activated saver, Rs.")
    ws["A2"].font = NOTE
    summ = read("mc_summary.csv")
    pct = [r for r in summ if r.get("p50", "") != ""]
    prob = [r for r in summ if r.get("p50", "") == ""]
    r = table(ws, 4, "Outputs: percentiles across draws (base_case = the deterministic middle case)", pct,
              ["variant", "metric", "base_case", "mean", "p5", "p10", "p25", "p50", "p75", "p90", "p95"],
              fmts={c: "#,##0.00" for c in ("base_case", "mean", "p5", "p10", "p25", "p50", "p75", "p90", "p95")},
              widths=[13, 12, 11, 11, 11, 11, 11, 11, 11, 11, 11])
    r = table(ws, r, "Probabilities (share of draws)", prob, ["variant", "metric", "mean"], fmts={"mean": "0.0%"})
    drv = sorted(read("mc_drivers.csv"), key=lambda x: (x["variant"], -float(x["variance_share"])))
    r = table(ws, r, "Drivers of net value per saver (rank regression; variance_share sums to 100% per variant)", drv,
              ["variant", "input", "spearman", "variance_share", "rank_model_r2"], fmts={"spearman": "0.00", "variance_share": "0.0%", "rank_model_r2": "0.00"})
    r = table(ws, r, "What moves the odds: P(net > 0) with the driver in its worst / best fifth", read("mc_quintiles.csv"),
              ["variant", "input", "p_net_pos_worst_fifth", "p_net_pos_best_fifth", "p_ratio_ge3_best_fifth"],
              fmts={c: "0.0%" for c in ("p_net_pos_worst_fifth", "p_net_pos_best_fifth", "p_ratio_ge3_best_fifth")})
    table(ws, r, "Histogram of net value per saver (Rs; 1st to 99th percentile)", read("mc_hist.csv"), ["variant", "lo", "hi", "share"],
          fmts={"lo": "#,##0", "hi": "#,##0", "share": "0.0%"})
    ws.freeze_panes = "A4"

    # ------------------------------------------------ sensitivity on net value
    ws = wb.create_sheet("Sensitivity (net)")
    ws["A1"] = "Sensitivity on net value per saver: one input at a time, and debit cost x retention"
    ws["A1"].font = TITLE
    ws["A2"] = ("Each input moved to the 10th and 90th percentile of its distribution, all others at the middle case. "
                "Break-even = the value at which the average saver's net value reaches zero with all else at the middle case "
                "(blank = no value in a wide range gets there). Grid: the whole retention curve, both halves, at one percentile.")
    ws["A2"].font = NOTE
    r = table(ws, 4, "One at a time (sorted by swing)", read("sens_oat.csv"),
              ["input", "p10_value", "base_value", "p90_value", "net_at_p10", "net_at_base", "net_at_p90", "swing", "breakeven_value"],
              fmts={c: "#,##0" for c in ("net_at_p10", "net_at_base", "net_at_p90", "swing")}, widths=[20, 11, 11, 11, 11, 11, 11, 10, 14])
    table(ws, r, "Net value per saver: debit cost (rows) x retention percentile", read("sens_grid.csv"),
          ["cost_debit", "retention_percentile", "ret_m12_gig", "net", "ratio"], fmts={"net": "#,##0", "ratio": "0.00", "ret_m12_gig": "0%"})
    return ["Assumption register", "Monte Carlo", "Sensitivity (net)"]


def kv(rows, key="metric", val="value"):
    return [{"item": r[key], "value": r[val]} for r in rows]


def build_redesign(wb):
    """Tabs for the work after the first verdict: the break-even minimum, cross-sell and one customer's
    value, the redesign (shopkeepers), the salaried version, and every customer group on one yardstick.
    All values written by the model scripts named on each tab; nothing recomputed here."""
    out = []

    ws = wb.create_sheet("Break-even minimum")
    ws["A1"] = "The flip: what daily minimum makes one saver break even?"
    ws["A1"].font = TITLE
    ws["A2"] = ("Written by model/flip.py. Every saver puts in exactly the minimum per successful debit; everything else at the evidence "
                "middle, so the savers a higher minimum would lose are NOT counted. Flat fee steps from Rs 5.90 to Rs 17.70 above Rs 1,000 a debit.")
    ws["A2"].font = NOTE
    r = table(ws, 4, "Break-even amount per case (first Rs 50 step with net value >= 0)", read("flip_breakeven.csv"),
              ["case", "design", "breakeven_amount", "monthly_gig", "share_income_gig", "best_net_at_or_below_1000"],
              fmts={"breakeven_amount": "#,##0", "monthly_gig": "#,##0", "share_income_gig": "0%", "best_net_at_or_below_1000": "#,##0"},
              widths=[30, 9, 16, 14, 16, 22])
    table(ws, r, "Monte Carlo by amount (2,000 draws each): share of futures with net value > 0", read("flip_mc.csv"),
          ["amount", "p_pos", "p_pos_flat", "p_pos_percent", "p50"], fmts={"p_pos": "0.0%", "p_pos_flat": "0.0%", "p_pos_percent": "0.0%", "p50": "#,##0"})
    out.append(ws.title)

    ws = wb.create_sheet("Cross-sell and LTV")
    ws["A1"] = "Everything we could sell, and what one customer is worth with it"
    ws["A1"].font = TITLE
    ws["A2"] = ("Cross-sell: take rate x ticket x yearly take-up per ACTIVE saver (take-up is the researchers' estimate; research/xs_*.csv). "
                "LTV: model/ltv.py, per customer won, 5 years discounted; cross-sell from month 1, before selling costs; replaces the model's own lines.")
    ws["A2"].font = NOTE
    r = table(ws, 4, "Cross-sell map (Rs per active saver a year)", read("crosssell_map.csv"),
              ["product", "group", "rs_per_active_saver_year_low", "mid", "high", "take_rate_mid", "verdict", "evidence"],
              fmts={"rs_per_active_saver_year_low": "#,##0.0", "mid": "#,##0.0", "high": "#,##0.0"}, widths=[44, 12, 12, 10, 10, 26, 34, 30])
    table(ws, r, "One customer, case by case (Rs per customer won)", read("ltv_cases.csv"),
          ["case", "cac", "savings_pv", "cross_sell_pv", "serve_pv", "ltv60", "net60", "payback_month"],
          fmts={c: "#,##0" for c in ("cac", "savings_pv", "cross_sell_pv", "serve_pv", "ltv60", "net60")})
    out.append(ws.title)

    for title, pre, note in (("Redesign - shopkeepers", "redesign",
                              "model/redesign.py and redesign_deck.py: shopkeepers save into debt funds (no fee per debit); partner lenders make Rs 50k-2.5 lakh unsecured loans; we earn 1-3% minus clawback."),
                             ("Redesign - salaried", "rs",
                              "model/redesign_salaried.py: salaried workers save from pay via employers; retention from staffing exits and tenure; loans Rs 50k-2.5 lakh; account moves with the worker unless stated.")):
        ws = wb.create_sheet(title)
        ws["A1"] = title.replace("Redesign - ", "The redesign: ")
        ws["A1"].font = TITLE
        ws["A2"] = note + " Futures = Monte Carlo draws; 'pays back' = net value per customer won > 0 within 5 years."
        ws["A2"].font = NOTE
        dist = read(f"{pre}_distributions.csv")
        r = table(ws, 4, "Inputs this version adds or changes", dist, ["input", "dist", "low", "mode", "high", "basis", "evidence"],
                  widths=[26, 11, 12, 12, 12, 60, 30])
        r = table(ws, r, "Each step: share of futures that pay back", read(f"{pre}_variants.csv"), ["case", "p_net_pos", "net_p50", "p_payback_le_36"],
                  fmts={"p_net_pos": "0.0%", "net_p50": "#,##0", "p_payback_le_36": "0.0%"})
        r = table(ws, r, "Grid: fixed cost to win x fixed loans a year per proven saver", read(f"{pre}_grid.csv"),
                  ["cac", "loans_year", "p_net_pos", "net_p50"], fmts={"cac": "#,##0", "loans_year": "0%", "p_net_pos": "0.0%", "net_p50": "#,##0"})
        r = table(ws, r, "Target case and today's evidence (Rs per customer won, 5 years discounted)", kv(read(f"{pre}_deck.csv")), ["item", "value"],
                  fmts={"value": "#,##0.00"})
        table(ws, r, "Drivers of the outcome (rank regression)", read(f"{pre}_drivers.csv")[:12], ["input", "variance_share"], fmts={"variance_share": "0.0%"})
        out.append(ws.title)

    ws = wb.create_sheet("Lending ourselves")
    ws["A1"] = "Lending ourselves: what one loan and one customer earn, by how well borrowers repay"
    ws["A1"].font = TITLE
    ws["A2"] = ("model/lend_self.py. Losses and running costs are shares of the loan balance a year: savers safer 3% + 3%; industry steady state "
                "6% + 5% (ICRA); industry today 9% + 8% (KreditBee, Kissht, Stashfin, Groww). Cost of borrowing 14%. Winning customers counted separately.")
    ws["A2"].font = NOTE
    r = table(ws, 4, "One loan", read("lend_loan.csv"), ["segment", "loan", "months", "rate", "case", "interest", "fee", "funding", "losses", "running", "net", "return_on_book", "referral_net"],
              fmts={c: "#,##0" for c in ("loan", "interest", "fee", "funding", "losses", "running", "net", "referral_net")} | {"rate": "0%", "return_on_book": "0.0%"},
              widths=[32, 10, 8, 7, 22, 10, 9, 10, 10, 10, 10, 12, 12])
    r = table(ws, r, "One customer (Rs, 5 years discounted, after the cost to win)", read("lend_customer.csv"),
              ["segment", "mode", "default", "takeup", "cac_case", "cac", "worth_before_cac", "net", "payback"],
              fmts={"takeup": "0%", "cac": "#,##0", "worth_before_cac": "#,##0", "net": "#,##0"})
    r = table(ws, r, "One shopkeeper worked through (1 in 5 borrowing a year)", read("lend_target.csv"),
              ["case", "cac", "savings", "loans", "other", "serve", "net", "payback"], fmts={c: "#,##0" for c in ("cac", "savings", "loans", "other", "serve", "net")})
    table(ws, r, "Customers won a year to cover an NBFC's fixed running cost", read("lend_scale.csv"),
          ["segment", "case", "takeup", "net_per_customer", "customers_a_year_for_3cr", "customers_a_year_for_5cr", "customers_a_year_for_8cr"],
          fmts={"net_per_customer": "#,##0", "customers_a_year_for_3cr": "#,##0", "customers_a_year_for_5cr": "#,##0", "customers_a_year_for_8cr": "#,##0"})
    out.append(ws.title)

    ws = wb.create_sheet("Lending Monte Carlo")
    ws["A1"] = "Lending ourselves: 10,000 futures per customer group, every input drawn from its evidence range"
    ws["A1"].font = TITLE
    ws["A2"] = "model/lend_mc.py. Value per customer won, five years discounted, after the cost to win. Lending inputs triangular and independent; customer inputs from the first design's evidence ranges."
    ws["A2"].font = NOTE
    r = table(ws, 4, "Lending inputs", read("lmc_distributions.csv"), ["input", "dist", "low", "mode", "high", "basis", "evidence"], widths=[24, 11, 12, 12, 12, 50, 40])
    r = table(ws, r, "Results", read("lmc_summary.csv"), ["segment", "p_pays_back", "p_payback_le_24", "p_pays_back_if_losses_le_4pct", "p_pays_back_if_losses_ge_7pct", "p5", "p50", "p95", "mean",
                                                           "p_pays_back_cac_100", "p_pays_back_cac_250", "p_pays_back_cac_500"],
              fmts={c: "0.0%" for c in ("p_pays_back", "p_payback_le_24", "p_pays_back_if_losses_le_4pct", "p_pays_back_if_losses_ge_7pct", "p_pays_back_cac_100", "p_pays_back_cac_250", "p_pays_back_cac_500")}
              | {c: "#,##0" for c in ("p5", "p50", "p95", "mean")})
    r = table(ws, r, "Salaried workers via employers: value per worker by rate and losses (Rs; 1 in 5 borrowing)", read("lend_salaried.csv"),
              ["cac_case", "rate", "case", "loan_net", "net", "payback"], fmts={"rate": "0%", "loan_net": "#,##0", "net": "#,##0"})
    r = table(ws, r, "Drivers (share of the spread)", [d for d in read("lmc_drivers.csv") if float(d["variance_share"]) >= 0.01],
              ["segment", "input", "variance_share", "spearman"], fmts={"variance_share": "0.0%", "spearman": "0.00"})
    r = table(ws, r, "Funnel: of 100 shopkeepers won", read("lend_funnel.csv"), ["takeup", "won", "saving_m6", "saving_m12", "loans_5y_per_100"],
              fmts={"takeup": "0%", "saving_m6": "0", "saving_m12": "0", "loans_5y_per_100": "0"})
    table(ws, r, "The prize (Rs crore; 1 in 5 borrowing; value after an NBFC running cost of Rs 5 crore)", read("lend_prize.csv"),
          ["won_a_year", "lent_a_year_cr", "avg_book_cr", "value_a_year_median_cr", "value_a_year_savers_safer_cr", "equity_colend_cr", "equity_own_book_cr"],
          fmts={c: "#,##0" for c in ("won_a_year", "lent_a_year_cr", "avg_book_cr", "value_a_year_median_cr", "value_a_year_savers_safer_cr", "equity_colend_cr", "equity_own_book_cr")})
    out.append(ws.title)

    ws = wb.create_sheet("Customer groups")
    ws["A1"] = "Which customer: the most we could pay to win one and still break even"
    ws["A1"].font = TITLE
    ws["A2"] = ("model/segment_value.py: the same redesign for every group; only loan size, how long people stay and whether loans continue "
                "after they stop saving differ. Evidence: research/segments.csv (SG rows) and segments_summary.md.")
    ws["A2"].font = NOTE
    table(ws, 4, "Customer groups (Rs per customer won)", read("segment_value.csv"),
          ["segment", "loan_size", "still_m12", "keep_borrowing", "max_cac_10", "max_cac_20", "cac_evidence", "evidence"],
          fmts={"loan_size": "#,##0", "still_m12": "0%", "max_cac_10": "#,##0", "max_cac_20": "#,##0", "cac_evidence": "#,##0"},
          widths=[40, 11, 10, 14, 12, 12, 12, 60])
    out.append(ws.title)
    return out
